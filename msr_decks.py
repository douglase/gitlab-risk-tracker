"""Generate per-product "Top 5 Risks" MSR slide decks from dashboard data.

For every product label found on open scored risks, fills the standard
monthly-status-report template (``templates/msr_top5.pptx`` — a single
slide with a 9-column table: Title, C, L, Score, Risk Type, Risk
Description, Trend, Notes, Mitigation Plan) with that product's five
highest-scored open risks and writes ``top5_<product>.pptx`` into the
output directory. When LibreOffice is available a print-ready PDF is
produced alongside each deck via ``scripts/pptx_to_pdf.sh`` (which also
handles the Aptos font substitution so the table doesn't overflow the
slide); otherwise decks are PPTX-only and a note is printed.

The Trend column mirrors the dashboard's 30-day movement section:
"New" (first seen within the window), then escalated, then
de-escalated, else steady.

SPDX-License-Identifier: GPL-3.0-or-later
Copyright (C) 2026 Ewan Douglas and contributors
"""

from __future__ import annotations

import copy
import re
import shutil
import subprocess
import sys
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn

TREND_NEW = "New"
TREND_UP = "↑"
TREND_DOWN = "↓"
TREND_FLAT = "→"

_SLUG_RE = re.compile(r"[^A-Za-z0-9_.-]+")


def trend_by_iid(mv: dict) -> dict:
    """Map iid -> trend glyph from a ``build.movement()`` result.

    Precedence: New beats escalated beats de-escalated; anything not in
    the movement lists is steady.
    """
    trends: dict = {}
    for r in mv.get("deescalated", []):
        trends[r["iid"]] = TREND_DOWN
    for r in mv.get("escalated", []):
        trends[r["iid"]] = TREND_UP
    for r in mv.get("new", []):
        trends[r["iid"]] = TREND_NEW
    return trends


# Inline markdown that must not leak into slide text as literal symbols.
_INLINE_MD: list[tuple[re.Pattern, str]] = [
    (re.compile(r"\[([^\]]+)\]\([^)]+\)"), r"\1"),          # [text](url)
    (re.compile(r"\*\*(.+?)\*\*"), r"\1"),                   # **bold**
    (re.compile(r"__(.+?)__"), r"\1"),                       # __bold__
    (re.compile(r"(?<![\w*])\*([^*\s][^*]*?)\*(?![\w*])"), r"\1"),  # *italic*
    (re.compile(r"(?<![\w_])_([^_\s][^_]*?)_(?![\w_])"), r"\1"),    # _italic_
    (re.compile(r"`([^`]*)`"), r"\1"),                       # `code`
]


def _strip_inline_md(text: str) -> str:
    for pattern, repl in _INLINE_MD:
        text = pattern.sub(repl, text)
    return text


def _section_lines(raw_md: str) -> list[str]:
    """Flatten a raw markdown section into slide-cell lines.

    Markdown bullets become the MSR house style ("-Do the thing." on its
    own line); soft-wrapped prose lines are joined back into single
    paragraphs; inline emphasis/link markup is stripped to plain text.
    """
    lines: list[str] = []
    prose: list[str] = []

    def flush() -> None:
        if prose:
            lines.append(_strip_inline_md(" ".join(prose)))
            prose.clear()

    for line in (raw_md or "").splitlines():
        line = line.strip()
        if not line:
            flush()
            continue
        m = re.match(r"^[-*+]\s+(.*)", line)
        if m:
            flush()
            lines.append(_strip_inline_md(f"-{m.group(1)}"))
        else:
            prose.append(line)
    flush()
    return lines


def _set_lines(text_frame, lines: list[str]) -> None:
    """Replace a text frame's content with one paragraph per line,
    cloning the first paragraph's formatting (the template holds a
    single placeholder run per cell)."""
    lines = lines or [""]
    p_elems = text_frame._txBody.findall(qn("a:p"))
    first = p_elems[0]
    for extra in p_elems[1:]:
        text_frame._txBody.remove(extra)
    proto = copy.deepcopy(first)

    def set_text(p_elem, text: str) -> None:
        r = p_elem.find(qn("a:r"))
        if r is None:  # placeholder cell with no run — plain text fallback
            r = copy.deepcopy(proto.find(qn("a:r")))
            if r is None:
                return
            p_elem.append(r)
        t = r.find(qn("a:t"))
        t.text = text

    set_text(first, lines[0])
    for line in lines[1:]:
        p = copy.deepcopy(proto)
        set_text(p, line)
        text_frame._txBody.append(p)


def _fill_deck(template: Path, out_pptx: Path, product: str,
               risks: list[dict], trends: dict) -> None:
    prs = Presentation(str(template))
    slide = prs.slides[0]
    table = next(sh for sh in slide.shapes if sh.has_table).table

    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.startswith("Top 5"):
            _set_lines(sh.text_frame, [f"Top 5 {product} Risks"])

    data_rows = list(table.rows)[1:]
    for row, it in zip(data_rows, risks):
        cells = row.cells
        secs = it.get("sections", {})
        _set_lines(cells[0].text_frame, [it["display_title"] or it["title"]])
        _set_lines(cells[1].text_frame, [str(it["consequence"])])
        _set_lines(cells[2].text_frame, [str(it["likelihood"])])
        _set_lines(cells[3].text_frame, [str(it["consequence"] * it["likelihood"])])
        _set_lines(cells[4].text_frame, [", ".join(it["risk_types"]) or "—"])
        _set_lines(cells[5].text_frame, _section_lines(secs.get("risk_description", "")))
        _set_lines(cells[6].text_frame, [trends.get(it["iid"], TREND_FLAT)])
        _set_lines(cells[7].text_frame, _section_lines(secs.get("notes", "")))
        _set_lines(cells[8].text_frame, _section_lines(secs.get("mitigation_plan", "")))
    for row in data_rows[len(risks):]:
        for cell in row.cells:
            _set_lines(cell.text_frame, [""])

    out_pptx.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_pptx))


def generate_msr_decks(items: list[dict], mv: dict, out_dir: Path,
                       template: Path, make_pdf: bool | None = None) -> list[dict]:
    """Write one Top-5 deck per product; return link metadata for the
    dashboard: [{product, count, pptx, pdf}] with paths relative to
    ``out_dir``'s parent (the public/ root)."""
    scored_open = [
        it for it in items
        if it["state"] != "closed"
        and it["consequence"] is not None and it["likelihood"] is not None
        and 1 <= it["consequence"] <= 5 and 1 <= it["likelihood"] <= 5
    ]
    products = sorted({p for it in scored_open for p in it["products"]})
    if not products:
        return []

    if make_pdf is None:
        make_pdf = shutil.which("soffice") is not None
    if not make_pdf:
        print("msr_decks: LibreOffice (soffice) not found — writing PPTX only.",
              file=sys.stderr)
    convert = Path(__file__).parent / "scripts" / "pptx_to_pdf.sh"

    trends = trend_by_iid(mv)
    decks: list[dict] = []
    for product in products:
        top5 = sorted(
            (it for it in scored_open if product in it["products"]),
            key=lambda it: (-(it["consequence"] * it["likelihood"]),
                            -it["consequence"], -it["likelihood"]),
        )[:5]
        slug = _SLUG_RE.sub("_", product).strip("_")
        pptx_path = out_dir / f"top5_{slug}.pptx"
        _fill_deck(template, pptx_path, product, top5, trends)
        pdf_rel = None
        if make_pdf:
            res = subprocess.run(
                ["bash", str(convert), str(pptx_path), str(out_dir)],
                capture_output=True, text=True, timeout=300,
            )
            if res.returncode == 0 and pptx_path.with_suffix(".pdf").exists():
                pdf_rel = f"{out_dir.name}/{pptx_path.with_suffix('.pdf').name}"
            else:
                print(f"msr_decks: PDF conversion failed for {pptx_path.name}:\n"
                      f"{res.stderr.strip()}", file=sys.stderr)
        decks.append({
            "product": product,
            "count": len(top5),
            "pptx": f"{out_dir.name}/{pptx_path.name}",
            "pdf": pdf_rel,
        })
    return decks
